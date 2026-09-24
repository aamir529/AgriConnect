import os
import sys
import base64
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ASSETS_DIR = os.path.join(BASE_DIR, 'static', 'report_assets')
LOGO_PATH = os.path.join(ASSETS_DIR, 'header_logo_2.png')
HTML_OUTPUT = os.path.join(BASE_DIR, 'AgriConnect_Project_Progress_Seminar_Report.html')
PDF_OUTPUT = os.path.join(BASE_DIR, 'AgriConnect_Project_Progress_Seminar_Report.pdf')
DOCX_OUTPUT = os.path.join(BASE_DIR, 'AgriConnect_Project_Progress_Seminar_Report.docx')
MD_OUTPUT = os.path.join(BASE_DIR, 'AgriConnect_Project_Progress_Seminar_Report.md')

# Read logo as base64 for standalone HTML
with open(LOGO_PATH, 'rb') as f:
    logo_base64 = base64.b64encode(f.read()).decode('utf-8')

# -------------------------------------------------------------
# 1. GENERATE HTML FOR HEADLESS EDGE PRINT-TO-PDF
# -------------------------------------------------------------
html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>AgriConnect - Project Progress Seminar Report</title>
<style>
  @page {{
    size: A4 portrait;
    margin: 0;
  }}
  * {{
    box-sizing: border-box;
    -webkit-print-color-adjust: exact;
    print-color-adjust: exact;
  }}
  body {{
    margin: 0;
    padding: 0;
    font-family: 'Times New Roman', Times, serif;
    color: #111;
    background-color: #fff;
    font-size: 10pt;
    line-height: 1.35;
  }}
  .page {{
    width: 210mm;
    height: 297mm;
    padding: 13mm 15mm 11mm 15mm;
    margin: 0 auto;
    page-break-after: always;
    page-break-inside: avoid;
    position: relative;
    overflow: hidden;
    background: #fff;
  }}
  .page:last-child {{
    page-break-after: avoid;
  }}

  /* Page 1 Header */
  .uni-header {{
    text-align: center;
    margin-bottom: 2px;
  }}
  .uni-header img {{
    max-width: 445px;
    height: auto;
    display: block;
    margin: 0 auto 3px auto;
  }}
  .dept-title {{
    font-size: 11.5pt;
    font-weight: bold;
    letter-spacing: 0.5px;
    margin: 3px 0 2px 0;
    text-transform: uppercase;
  }}
  .seminar-title-row {{
    font-size: 11pt;
    font-weight: bold;
    margin: 2px 0 2px 0;
    display: flex;
    align-items: center;
    justify-content: center;
    gap: 6px;
  }}
  .checkbox-box {{
    display: inline-block;
    width: 32px;
    height: 18px;
    border: 1.5px solid #222;
    border-radius: 3px;
    text-align: center;
    line-height: 16px;
    font-size: 9.5pt;
    font-weight: bold;
  }}
  .on-text {{
    font-size: 10pt;
    font-weight: bold;
    margin: 2px 0 2px 0;
  }}
  .project-title-box {{
    border-bottom: 1.2px solid #333;
    padding-bottom: 3px;
    margin-bottom: 5px;
    font-size: 10.8pt;
    font-weight: bold;
    color: #0b3d20;
    text-align: center;
  }}
  .prep-label {{
    text-align: center;
    font-style: italic;
    font-size: 9.5pt;
    margin: 1px 0 2px 0;
  }}

  /* Table on Page 1 */
  .student-table {{
    width: 100%;
    border-collapse: collapse;
    margin-bottom: 5px;
    font-size: 9pt;
  }}
  .student-table th, .student-table td {{
    border: 1px solid #333;
    padding: 3.2px 6px;
    text-align: center;
  }}
  .student-table th {{
    background-color: #f8f9fa;
    font-weight: bold;
    font-size: 8.5pt;
  }}
  .student-table td.left {{
    text-align: left;
  }}

  .guidance-section {{
    text-align: center;
    margin: 4px 0 4px 0;
    font-size: 9.5pt;
  }}
  .guidance-label {{
    font-style: italic;
    margin-bottom: 2px;
  }}
  .guidance-line {{
    border-bottom: 1px solid #444;
    width: 75%;
    margin: 0 auto;
    padding-bottom: 1px;
    font-weight: bold;
  }}
  .date-row {{
    margin: 3px 0 5px 0;
    font-size: 9.5pt;
    font-weight: bold;
  }}

  /* Framed Space for Writing */
  .frame-box {{
    border: 1.2px solid #222;
    padding: 8px 10px;
    position: relative;
    border-radius: 1px;
  }}
  .frame-title {{
    text-align: center;
    font-style: italic;
    text-decoration: underline;
    font-weight: bold;
    font-size: 9.5pt;
    margin-bottom: 5px;
  }}
  .sec-heading {{
    font-weight: bold;
    font-size: 10.5pt;
    color: #0d2b1d;
    margin: 4px 0 3px 0;
    text-decoration: underline;
  }}
  .body-p {{
    text-align: justify;
    margin: 0 0 5px 0;
    text-justify: inter-word;
    font-size: 9.1pt;
    line-height: 1.34;
  }}
  .bullet-list {{
    margin: 2px 0 5px 0;
    padding-left: 18px;
    font-size: 9.1pt;
    line-height: 1.32;
  }}
  .bullet-list li {{
    margin-bottom: 2px;
    text-align: justify;
  }}

  /* Page 2 Specific */
  .page2-frame {{
    height: 98%;
    border: 1.2px solid #222;
    padding: 10px 12px;
  }}
  .sub-sec {{
    margin-top: 4px;
    margin-bottom: 2px;
    font-weight: bold;
    font-size: 9.5pt;
    color: #1a3d2f;
  }}

  /* Page 3 Specific */
  .page3-frame {{
    height: 98%;
    border: 1.2px solid #222;
    padding: 10px 12px;
  }}
  .progress-sublist {{
    margin: 2px 0 4px 0;
    padding-left: 16px;
    list-style-type: lower-alpha;
    font-size: 8.8pt;
    line-height: 1.3;
  }}
  .progress-sublist li {{
    margin-bottom: 2.5px;
    text-align: justify;
  }}
  .pct-badge {{
    display: inline-block;
    background: #e8f5e9;
    border: 1.2px solid #2e7d32;
    color: #1b5e20;
    padding: 2px 8px;
    font-weight: bold;
    border-radius: 4px;
    font-size: 10pt;
    margin-right: 6px;
  }}

  /* Page 4 Specific */
  .page4-container {{
    height: 100%;
    padding: 6mm 4mm;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
  }}
  .sig-student-row {{
    display: flex;
    justify-content: space-between;
    align-items: flex-end;
    margin-bottom: 12px;
    font-size: 10pt;
    font-weight: normal;
  }}
  .sig-label {{
    white-space: nowrap;
    font-weight: 500;
  }}
  .sig-underline {{
    border-bottom: 1px dotted #333;
    flex-grow: 1;
    margin-left: 8px;
    height: 15px;
  }}
  .remarks-block {{
    margin-top: 8px;
    padding-top: 4px;
  }}
  .remarks-title {{
    font-weight: bold;
    font-size: 10.5pt;
    margin-bottom: 4px;
    letter-spacing: 0.2px;
  }}
  .ruled-lines {{
    width: 100%;
    margin-bottom: 8px;
  }}
  .ruled-line {{
    border-bottom: 1px solid #777;
    height: 19px;
    width: 100%;
  }}
  .sign-footer {{
    display: flex;
    justify-content: space-between;
    align-items: flex-end;
    font-size: 9.5pt;
    font-weight: bold;
    margin-top: 6px;
  }}
</style>
</head>
<body>

<!-- ==================== PAGE 1 ==================== -->
<div class="page">
  <div class="uni-header">
    <img src="data:image/png;base64,{logo_base64}" alt="ARKA JAIN University Jharkhand">
    <div class="dept-title">DEPARTMENT OF COMPUTER SCIENCE ENGINEERING</div>
    <div class="seminar-title-row">
      <span>PROJECT PROGRESS SEMINAR -</span>
      <div class="checkbox-box">II</div>
    </div>
    <div class="on-text">ON</div>
    <div class="project-title-box">
      AgriConnect: Transparent Digital Agricultural Supply Chain Platform with Assisted-Digital Access
    </div>
  </div>

  <div class="prep-label">Prepared by</div>
  <table class="student-table">
    <thead>
      <tr>
        <th style="width: 8%;">Sl. No.</th>
        <th style="width: 32%;">NAME</th>
        <th style="width: 20%;">Enrolment No.</th>
        <th style="width: 18%;">Class Roll No.</th>
        <th style="width: 22%;">Role in Project</th>
      </tr>
    </thead>
    <tbody>
      <tr>
        <td>1.</td>
        <td class="left"><strong>Aamir Alam</strong></td>
        <td>AJU/220541</td>
        <td>CS-2022-34</td>
        <td><strong>Leader</strong></td>
      </tr>
      <tr>
        <td>2.</td>
        <td class="left">[Student 2 Name]</td>
        <td>[Enrolment No.]</td>
        <td>[Class Roll No.]</td>
        <td>Frontend & UI Architecture</td>
      </tr>
      <tr>
        <td>3.</td>
        <td class="left">[Student 3 Name]</td>
        <td>[Enrolment No.]</td>
        <td>[Class Roll No.]</td>
        <td>Backend & Database Systems</td>
      </tr>
      <tr>
        <td>4.</td>
        <td class="left">[Student 4 Name]</td>
        <td>[Enrolment No.]</td>
        <td>[Class Roll No.]</td>
        <td>Logistics & Telephony Engine</td>
      </tr>
      <tr>
        <td>5.</td>
        <td class="left">[Student 5 Name]</td>
        <td>[Enrolment No.]</td>
        <td>[Class Roll No.]</td>
        <td>ML Analytics & QA Testing</td>
      </tr>
    </tbody>
  </table>

  <div class="guidance-section">
    <div class="guidance-label">Under the guidance of</div>
    <div class="guidance-line">[Project Guide / Supervisor Name, Designation, Dept. of CSE]</div>
  </div>

  <div class="date-row">
    Date : <u>&nbsp;&nbsp;15 / 09 / 2026&nbsp;&nbsp;</u>
  </div>

  <div class="frame-box" style="height: 122mm;">
    <div class="frame-title">Space for Writing</div>
    
    <div class="sec-heading">Scope of the Project:</div>
    <p class="body-p">
      <strong>Problem Addressed:</strong> Traditional Indian agricultural supply chains suffer from a multi-tiered intermediary hierarchy (Local Brokers &rarr; Commission Agents &rarr; Mandi Wholesalers &rarr; Distributors &rarr; Urban Retailers). Consequently, smallholder farmers receive an exploitative 20%–30% share of the final consumer expenditure. Concurrently, smallholders face severe digital literacy and smartphone access hurdles that prevent adoption of conventional e-commerce apps.
    </p>
    <p class="body-p">
      <strong>Target Users:</strong> (1) Rural smallholder farmers; (2) Village Digital Facilitators (VDFs / CSC kiosk operators); (3) Urban residential consumers and Housing Societies (RWAs); (4) Local transporters/van fleet drivers; (5) Platform agricultural cooperative administrators.
    </p>
    <p class="body-p">
      <strong>Main Purpose & Major Functionality:</strong> AgriConnect ("Krishi Setu") disintermediates the exploitative chain through an assisted-digital ("phygital") kiosk and dual-tone IVR phone intake, automated price decomposition ("Where Does Your Rupee Go?"), community group buying pools, Leaflet.js farm-radius discovery, Agmarknet mandi market intelligence, two-party OTP dispatch verification, and Scikit-learn crop demand forecasting.
    </p>
    <p class="body-p">
      <strong>Expected Use of the System:</strong> Deployed across rural village clusters to aggregate produce at village kiosks, direct-route batches to urban societies, provide fair farmgate cash realization (+40% higher than local traders), and provide 15% savings to consumers.
    </p>

    <div class="sec-heading">About the Project:</div>
    <p class="body-p">
      <strong>Project Overview:</strong> AgriConnect is a production-engineered full-stack agricultural logistics and commerce platform. Built using Django 5.1 and Python 3.11+, it bridges digital exclusion in agriculture through community kiosks, simulated telephony, transparent financial margin distribution, and machine-learning driven market intelligence.
    </p>
  </div>
</div>

<!-- ==================== PAGE 2 ==================== -->
<div class="page">
  <div class="page2-frame">
    <div class="sec-heading" style="margin-top: 0;">About the Project: (Continued)</div>
    
    <div class="sub-sec">1. System Workflow & Operational Mechanism</div>
    <p class="body-p">
      AgriConnect replaces exploitative multi-tier middleman networks with a streamlined 4-stage assisted digital pipeline:
    </p>
    <ul class="bullet-list">
      <li><strong>Assisted Farmer Onboarding & Listing:</strong> Farmers with basic keypad feature phones interact with local Village Digital Facilitators (VDFs) or dial a toll-free IVR simulator. Facilitators inspect produce quality, log crop specifications (variety, weight, harvest date), and verify baseline cultivation costs.</li>
      <li><strong>Automated Fair Price Discovery:</strong> The platform cross-references real-time Agmarknet mandi benchmark feeds from <code>data.gov.in</code> and computes recommended farmgate price corridors (&ge; Modal Mandi Price &times; 1.2), eliminating distress selling.</li>
      <li><strong>Direct Consumer Marketplace & Group Pools:</strong> Urban households and residential societies order produce with clear margin transparency breakdowns. Society members join collective "Group Buying Pools" to trigger bulk volume discounts and consolidate delivery logistics.</li>
      <li><strong>Consolidated Logistics & Two-Party Verification:</strong> Aggregated harvests at VDF kiosks are claimed by local transporters via dispatch manifests. Delivery milestones are tracked via Leaflet GPS routes, and payments are settled upon two-party OTP verification.</li>
    </ul>

    <div class="sub-sec">2. Core Modules & Implemented Features</div>
    <ul class="bullet-list">
      <li><strong>Accounts & 5-Role RBAC:</strong> Granular role-based access control supporting Farmer, Facilitator, Consumer, Transporter, and Administrator with customized dashboards and session isolation.</li>
      <li><strong>Village Digital Facilitator Hub & IVR Engine:</strong> Assisted onboarding workflows, voice produce listing intake, and dual-tone multi-frequency (DTMF) simulated telephony engine with Web Audio and Hindi/English speech synthesis.</li>
      <li><strong>Direct Marketplace & Price Decomposition:</strong> Real-time margin calculator displaying farmer realization, freight, VDF commission, and platform fees ("Where Does Your Rupee Go?"). Features Leaflet.js farm cluster maps and QR-based Farm-to-Fork provenance tracking.</li>
      <li><strong>Order Management, Dispatch & Live Tracking:</strong> Session cart, checkout with mock UPI/COD, dispatch sheets, Leaflet waypoint GPS tracking, and two-tier quality dispute arbitration desks.</li>
      <li><strong>Market Intelligence & Scikit-Learn Demand Forecasting:</strong> Live Agmarknet mandi price comparisons across commodities and districts, paired with a trained Scikit-learn regression model for 7-day crop demand volume projections.</li>
      <li><strong>Value-Added Agri-Services:</strong> Digital Kisan Credit Passbook tracking harvest turnover and creditworthiness, AI Crop Doctor bio-advisory for disease diagnosis, and IoT cold-chain telemetry tracking transit temperature and humidity.</li>
    </ul>

    <div class="sub-sec">3. Technologies Used & Architectural Stack</div>
    <table class="student-table" style="margin-top: 4px; font-size: 8.8pt;">
      <thead>
        <tr>
          <th style="width: 22%;">Layer</th>
          <th style="width: 38%;">Technology Stack</th>
          <th style="width: 40%;">Architectural Role & Scope</th>
        </tr>
      </thead>
      <tbody>
        <tr>
          <td><strong>Frontend</strong></td>
          <td class="left">HTML5, CSS3, Bootstrap 5, JavaScript (ES6+), Leaflet.js, Chart.js</td>
          <td class="left">Responsive UI/UX, interactive map clustering, dynamic price decomposition graphs, and PWA offline capability.</td>
        </tr>
        <tr>
          <td><strong>Backend</strong></td>
          <td class="left">Python 3.11+, Django 5.1, Django REST Framework (DRF)</td>
          <td class="left">MVC architecture, secure authentication, transactional order pipelines, RESTful APIs, and business domain logic.</td>
        </tr>
        <tr>
          <td><strong>Database</strong></td>
          <td class="left">SQLite (Evaluation) / PostgreSQL (Production)</td>
          <td class="left">Relational schema, ACID transaction guarantees, index-optimized spatial and catalog queries.</td>
        </tr>
        <tr>
          <td><strong>Key Integrations</strong></td>
          <td class="left">data.gov.in Agmarknet API, OpenStreetMap, Scikit-learn, Web Audio IVR</td>
          <td class="left">Official APMC mandi feeds, geospatial routing, 7-day demand forecasting regressor, and browser telephony simulation.</td>
        </tr>
      </tbody>
    </table>
  </div>
</div>

<!-- ==================== PAGE 3 ==================== -->
<div class="page">
  <div class="page3-frame">
    <div class="sec-heading" style="margin-top: 0;">Narrative Summary of Project Status:</div>
    <p class="body-p">
      The AgriConnect platform has successfully transitioned from conceptual system blueprint to a fully implemented, production-tested software system. All primary architectural layers—encompassing the 5-role authentication engine, assisted-digital VDF kiosk, DTMF/IVR telephony simulator, direct consumer marketplace with dynamic margin decomposition, order fulfillment dispatch hub, Leaflet GPS live route tracker, Agmarknet mandi market intelligence integration, and Scikit-learn demand forecasting engine—are completely developed and operational. 
    </p>
    <p class="body-p">
      Comprehensive automated verification has been executed via <code>test_full_suite.py</code> across all 8 architectural modules, yielding <strong>48 passed checks out of 48 total assertions (100% pass rate, 0 defects)</strong>. The system is currently at the completed pre-deployment and evaluation stage, ready for real-world village cluster pilot trials.
    </p>

    <div class="sec-heading">Progress Summary:</div>
    <div style="font-weight: bold; font-size: 9.3pt; margin-top: 2px;">1. Completed tasks / Activities</div>
    <ol class="progress-sublist">
      <li><strong>5-Role Role-Based Access Control (RBAC) System:</strong> Developed custom user authentication and secure profile schemas catering to Farmers, Facilitators (VDFs), Consumers, Transporters, and Administrators with isolated dashboards and permission guards.</li>
      <li><strong>Assisted-Digital Kiosk & IVR Telephony Simulator:</strong> Engineered the Village Digital Facilitator portal alongside a browser-based dual-tone multi-frequency (DTMF) IVR and SMS notification engine for feature-phone and low-literacy farmers.</li>
      <li><strong>Direct Marketplace & Dynamic Price Decomposition:</strong> Implemented the consumer-facing fresh produce catalog featuring interactive "Where Does Your Rupee Go?" financial breakdown cards, community Group Buying pools, and Leaflet.js farm cluster maps.</li>
      <li><strong>End-to-End Order & Dispatch Logistics Pipeline:</strong> Built multi-item session cart, checkout, dispatch manifests, Leaflet GPS live tracking, two-party OTP delivery verification, and photo-supported dispute arbitration desks.</li>
      <li><strong>Market Intelligence & Scikit-Learn ML Forecasting:</strong> Integrated live <code>data.gov.in</code> Agmarknet mandi benchmark feeds, dynamic fair-price corridor computation, and trained Random Forest / Ridge ML models for 7-day crop demand prediction.</li>
    </ol>

    <div style="font-weight: bold; font-size: 9.3pt; margin-top: 4px;">2. Planned tasks / Activities</div>
    <ol class="progress-sublist">
      <li><strong>Live Rural Village Pilot Deployment:</strong> Execute field deployment across selected village panchayats in Jharkhand to evaluate real-world kiosk operations with active farmer and VDF cohorts.</li>
      <li><strong>Telecommunication Cloud Gateway Binding:</strong> Integrate live telecom SMS gateway (Twilio / SMSCountry) and IVR PRI cloud trunks to transition from browser simulation to live mobile network lines.</li>
      <li><strong>Production Cloud Infrastructure Migration:</strong> Transition from local SQLite evaluation environment to high-availability cloud architecture (PostgreSQL, Redis caching, Gunicorn/Nginx on AWS/GCP).</li>
      <li><strong>Regional Vernacular Voice NLP Expansion:</strong> Extend automated voice listing engine to support local spoken dialects (Hindi, Santhali, Mundari, Ho) via advanced ASR speech models.</li>
      <li><strong>Hardware IoT Sensor Telemetry Integration:</strong> Interface physical microcontroller hardware (ESP32/DHT22) for real-time transit vehicle temperature, humidity, and cold-chain logging.</li>
    </ol>

    <div class="sec-heading">Problems Faced (if any):</div>
    <p class="body-p" style="margin-bottom: 3px;">
      During system engineering, four critical technical and socio-technical challenges were identified and systematically resolved:
    </p>
    <ul class="bullet-list" style="font-size: 8.7pt;">
      <li><strong>Rural Digital Illiteracy & Device Constraints:</strong> Smallholder farmers often lack smartphones and data connectivity. <em>Resolution:</em> Conceived and engineered the "Phygital" Village Digital Facilitator (VDF) assisted model and dual-tone IVR phone intake.</li>
      <li><strong>Government Mandi API Schema Variance & Latency:</strong> data.gov.in Agmarknet endpoints occasionally exhibited variable response latency. <em>Resolution:</em> Implemented an intelligent asynchronous caching layer with verified fallback benchmark prices.</li>
      <li><strong>Mathematical Balancing of Dynamic Margins:</strong> Breaking down prices transparently across logistics, commissions, and farmer base values required strict validation. <em>Resolution:</em> Engineered real-time price decomposition algorithms in the product model.</li>
      <li><strong>Perishable Cold-Chain & Quality Verification:</strong> Transit damages risk multi-stakeholder disputes. <em>Resolution:</em> Developed two-party OTP delivery confirmation and photo-documented dispute arbitration workflows.</li>
    </ul>

    <div class="sec-heading" style="margin-top: 4px;">Project Completion Percentage:</div>
    <div style="margin-top: 3px; font-size: 9.2pt;">
      <span class="pct-badge">92% COMPLETED</span>
      <strong>Evaluation Rationale:</strong> 100% of all software requirements, database models, business logic controllers, interactive UI views, and the 48-unit automated test suite have been built and rigorously verified. The remaining 8% is designated for real-world telecom PRI line binding, physical IoT sensor hookup, and rural field pilot trials.
    </div>
  </div>
</div>

<!-- ==================== PAGE 4 ==================== -->
<div class="page">
  <div class="page4-container">
    <div style="margin-top: 2mm;">
      <div class="sig-student-row">
        <span class="sig-label">Signature of Student 1-</span>
        <div class="sig-underline"></div>
      </div>
      <div class="sig-student-row">
        <span class="sig-label">Signature of Student 2-</span>
        <div class="sig-underline"></div>
      </div>
      <div class="sig-student-row">
        <span class="sig-label">Signature of Student 3-</span>
        <div class="sig-underline"></div>
      </div>
      <div class="sig-student-row">
        <span class="sig-label">Signature of Student 4-</span>
        <div class="sig-underline"></div>
      </div>
      <div class="sig-student-row">
        <span class="sig-label">Signature of Student 5-</span>
        <div class="sig-underline"></div>
      </div>
    </div>

    <!-- Supervisor Remarks -->
    <div class="remarks-block">
      <div class="remarks-title">Project Supervisor’s Remarks:</div>
      <div class="ruled-lines">
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
      </div>
      <div class="sign-footer">
        <div>Date : __________________________</div>
        <div>Project Supervisor Signature</div>
      </div>
    </div>

    <!-- Coordinator Remarks -->
    <div class="remarks-block">
      <div class="remarks-title">Program Coordinator’s Remarks</div>
      <div class="ruled-lines">
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
      </div>
      <div class="sign-footer">
        <div>Date : __________________________</div>
        <div>Program Coordinator’s Signature</div>
      </div>
    </div>

    <!-- Dean Remarks -->
    <div class="remarks-block">
      <div class="remarks-title">Dean Remarks</div>
      <div class="ruled-lines">
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
        <div class="ruled-line"></div>
      </div>
      <div class="sign-footer">
        <div>Date : __________________________</div>
        <div>Dean’s Signature:</div>
      </div>
    </div>
  </div>
</div>

</body>
</html>
"""

with open(HTML_OUTPUT, 'w', encoding='utf-8') as f:
    f.write(html_content)
print(f"Generated HTML report: {HTML_OUTPUT}")

# -------------------------------------------------------------
# 2. GENERATE PDF USING HEADLESS EDGE
# -------------------------------------------------------------
edge_cmd = [
    r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',
    '--headless=new',
    '--disable-gpu',
    f'--print-to-pdf={PDF_OUTPUT}',
    '--no-pdf-header-footer',
    HTML_OUTPUT
]
print("Running headless Edge to render PDF...")
proc = subprocess.run(edge_cmd, capture_output=True, text=True)
if proc.returncode != 0:
    print(f"Edge error: {proc.stderr}")
else:
    print(f"Generated PDF report: {PDF_OUTPUT} (Size: {os.path.getsize(PDF_OUTPUT)} bytes)")

# -------------------------------------------------------------
# 3. GENERATE EDITABLE DOCX DOCUMENT USING PYTHON-DOCX
# -------------------------------------------------------------
doc = Document()

# Set standard A4 page size and margins
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'<w:tcBorders %s/>' % nsdecls('w'))
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = parse_xml(r'<%s %s w:w="%s" w:val="%s" w:space="0" w:color="%s"/>' %
                                (tag, nsdecls('w'), edge_data.get('sz', '4'), edge_data.get('val', 'single'), edge_data.get('color', 'auto')))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), fill_hex))
    tcPr.append(shd)

def add_p(doc, text="", bold=False, italic=False, font_size=10, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
    return p

# --- PAGE 1 ---
p_logo = doc.add_paragraph()
p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_logo.paragraph_format.space_after = Pt(2)
run_logo = p_logo.add_run()
run_logo.add_picture(LOGO_PATH, width=Inches(4.6))

add_p(doc, "DEPARTMENT OF COMPUTER SCIENCE ENGINEERING", bold=True, font_size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
add_p(doc, "PROJECT PROGRESS SEMINAR - [ II ]", bold=True, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
add_p(doc, "ON", bold=True, font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

p_title = add_p(doc, "AgriConnect: Transparent Digital Agricultural Supply Chain Platform with Assisted-Digital Access", bold=True, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
p_title.runs[0].font.color.rgb = RGBColor(11, 61, 32)

add_p(doc, "Prepared by", italic=True, font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)

# Student Table
table1 = doc.add_table(rows=6, cols=5)
table1.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Sl. No.", "NAME", "Enrolment No.", "Class Roll No.", "Role in Project"]
widths = [Inches(0.6), Inches(2.3), Inches(1.3), Inches(1.2), Inches(1.6)]

for col_idx, text in enumerate(headers):
    cell = table1.cell(0, col_idx)
    cell.width = widths[col_idx]
    set_cell_background(cell, "F2F2F2")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(8.5)

student_data = [
    ("1.", "Aamir Alam", "AJU/220541", "CS-2022-34", "Leader"),
    ("2.", "[Student 2 Name]", "[Enrolment No.]", "[Class Roll No.]", "Frontend & UI Architecture"),
    ("3.", "[Student 3 Name]", "[Enrolment No.]", "[Class Roll No.]", "Backend & Database Systems"),
    ("4.", "[Student 4 Name]", "[Enrolment No.]", "[Class Roll No.]", "Logistics & Telephony Engine"),
    ("5.", "[Student 5 Name]", "[Enrolment No.]", "[Class Roll No.]", "ML Analytics & QA Testing"),
]

for row_idx, row_data in enumerate(student_data, start=1):
    for col_idx, val in enumerate(row_data):
        cell = table1.cell(row_idx, col_idx)
        cell.width = widths[col_idx]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 1 else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        run = p.add_run(val)
        if row_idx == 1:
            run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(8.5)

for row in table1.rows:
    for cell in row.cells:
        set_cell_border(cell, top=dict(sz=4, val='single', color='333333'),
                              bottom=dict(sz=4, val='single', color='333333'),
                              left=dict(sz=4, val='single', color='333333'),
                              right=dict(sz=4, val='single', color='333333'))

add_p(doc, "Under the guidance of", italic=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
add_p(doc, "[Project Guide / Supervisor Name, Designation, Department of CSE]", bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
add_p(doc, "Date : 15 / 09 / 2026", bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)

# Framed Box: Space for Writing (Page 1)
table_box1 = doc.add_table(rows=1, cols=1)
table_box1.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_box1 = table_box1.cell(0, 0)
cell_box1.width = Inches(7.0)
set_cell_border(cell_box1, top=dict(sz=8, val='single', color='222222'),
                           bottom=dict(sz=8, val='single', color='222222'),
                           left=dict(sz=8, val='single', color='222222'),
                           right=dict(sz=8, val='single', color='222222'))

p_box_title = cell_box1.paragraphs[0]
p_box_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_box_title.paragraph_format.space_after = Pt(3)
r = p_box_title.add_run("Space for Writing")
r.italic = True
r.bold = True
r.underline = True
r.font.name = 'Times New Roman'
r.font.size = Pt(9.5)

def add_box_p(cell, text="", bold_prefix="", font_size=9, space_after=3):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(font_size)
    if text:
        r_txt = p.add_run(text)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(font_size)
    return p

add_box_p(cell_box1, "", bold_prefix="Scope of the Project:", font_size=10, space_after=2)
add_box_p(cell_box1, " Traditional Indian agricultural supply chains involve a fragmented multi-tier middleman hierarchy (Village Brokers → Commission Agents → Mandi Wholesalers → Sub-wholesalers → Retailers). Consequently, smallholder farmers receive an exploitative 20%–30% share of consumer retail expenditure. Furthermore, rural smallholders face severe digital literacy and smartphone barriers, preventing adoption of standard e-commerce platforms.", bold_prefix="Problem Addressed:")
add_box_p(cell_box1, " (1) Rural smallholder farmers; (2) Village Digital Facilitators (VDFs / rural kiosk operators); (3) Urban residential consumers and Housing Societies (RWAs); (4) Local transporters and tempo drivers; (5) Platform agricultural administrators and dispute mediators.", bold_prefix="Target Users:")
add_box_p(cell_box1, " Disintermediate exploitative supply chains through an assisted-digital ('phygital') kiosk and dual-tone IVR phone intake, automated price decomposition ('Where Does Your Rupee Go?'), community group buying pools, Leaflet.js farm-radius discovery, Agmarknet mandi market intelligence, two-party OTP dispatch verification, and Scikit-learn crop demand forecasting.", bold_prefix="Main Purpose & Major Functionality:")
add_box_p(cell_box1, " Deployed across rural village clusters to aggregate produce at village kiosks, direct-route batches to urban societies, provide fair farmgate cash realization (+40% higher than local traders), and provide 15% savings to consumers.", bold_prefix="Expected Use of the System:")
add_box_p(cell_box1, "", bold_prefix="About the Project:", font_size=10, space_after=2)
add_box_p(cell_box1, " AgriConnect is a production-engineered full-stack agricultural logistics and commerce platform built with Django 5.1 and Python 3.11+. It bridges digital exclusion in agriculture through community kiosks, simulated telephony, transparent financial margin distribution, and machine-learning driven market intelligence.", bold_prefix="Project Overview:")

# --- PAGE 2 ---
doc.add_page_break()

table_box2 = doc.add_table(rows=1, cols=1)
table_box2.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_box2 = table_box2.cell(0, 0)
cell_box2.width = Inches(7.0)
set_cell_border(cell_box2, top=dict(sz=8, val='single', color='222222'),
                           bottom=dict(sz=8, val='single', color='222222'),
                           left=dict(sz=8, val='single', color='222222'),
                           right=dict(sz=8, val='single', color='222222'))

p2_title = cell_box2.paragraphs[0]
p2_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
p2_title.paragraph_format.space_after = Pt(3)
r2 = p2_title.add_run("About the Project: (Continued)")
r2.bold = True
r2.font.name = 'Times New Roman'
r2.font.size = Pt(10.5)

add_box_p(cell_box2, "", bold_prefix="1. System Workflow & Operational Mechanism", font_size=9.5, space_after=2)
add_box_p(cell_box2, "AgriConnect replaces multi-tier middlemen with a streamlined 4-stage assisted digital pipeline:", font_size=9, space_after=2)
add_box_p(cell_box2, " Farmers with basic keypad feature phones interact with local Village Digital Facilitators (VDFs) or dial a toll-free IVR simulator. Facilitators inspect produce quality, log crop specifications (variety, weight, harvest date), and verify baseline cultivation costs.", bold_prefix="• Assisted Onboarding & Listing:")
add_box_p(cell_box2, " The platform cross-references real-time Agmarknet mandi benchmark feeds from data.gov.in and computes recommended farmgate price corridors (≥ Modal Mandi Price × 1.2), preventing distress selling during glut periods.", bold_prefix="• Automated Fair Price Discovery:")
add_box_p(cell_box2, " Urban households and residential societies order produce with clear margin transparency breakdowns. Society members join collective 'Group Buying Pools' to trigger bulk volume discounts and consolidate delivery logistics.", bold_prefix="• Direct Marketplace & Group Pools:")
add_box_p(cell_box2, " Aggregated harvests at VDF kiosks are claimed by local transporters via dispatch manifests. Delivery milestones are tracked via Leaflet GPS routes, and payments are settled upon two-party OTP verification.", bold_prefix="• Consolidated Logistics & OTP Settlement:")

add_box_p(cell_box2, "", bold_prefix="2. Core Modules & Implemented Features", font_size=9.5, space_after=2)
add_box_p(cell_box2, " Granular role-based access control supporting Farmer, Facilitator, Consumer, Transporter, and Administrator with customized dashboards and session isolation.", bold_prefix="• Accounts & 5-Role RBAC:")
add_box_p(cell_box2, " Assisted onboarding workflows, voice produce listing intake, and dual-tone multi-frequency (DTMF) simulated telephony engine with Web Audio and Hindi/English speech synthesis.", bold_prefix="• Village Digital Facilitator Hub & IVR Engine:")
add_box_p(cell_box2, " Real-time margin calculator displaying farmer realization, freight, VDF commission, and platform fees ('Where Does Your Rupee Go?'). Features Leaflet.js farm cluster maps and QR-based Farm-to-Fork provenance tracking.", bold_prefix="• Direct Marketplace & Price Decomposition:")
add_box_p(cell_box2, " Session cart, checkout with mock UPI/COD, dispatch sheets, Leaflet waypoint GPS tracking, and two-tier quality dispute arbitration desks.", bold_prefix="• Order Management, Dispatch & Live Tracking:")
add_box_p(cell_box2, " Live Agmarknet mandi price comparisons across commodities and districts, paired with a trained Scikit-learn regression model for 7-day crop demand volume projections.", bold_prefix="• Market Intelligence & ML Demand Forecasting:")
add_box_p(cell_box2, " Digital Kisan Credit Passbook tracking harvest turnover, AI Crop Doctor bio-advisory for disease diagnosis, and IoT cold-chain telemetry tracking transit temperature and humidity.", bold_prefix="• Value-Added Agri-Services:")

add_box_p(cell_box2, "", bold_prefix="3. Technologies Used & Architectural Stack", font_size=9.5, space_after=2)

t_tech = cell_box2.add_table(rows=5, cols=3)
t_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
t_headers = ["Layer", "Technology Stack", "Architectural Role & Scope"]
t_widths = [Inches(1.2), Inches(2.3), Inches(3.2)]

for c_idx, h_txt in enumerate(t_headers):
    c = t_tech.cell(0, c_idx)
    c.width = t_widths[c_idx]
    set_cell_background(c, "F2F2F2")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    r = p.add_run(h_txt)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(8)

tech_rows = [
    ("Frontend", "HTML5, CSS3, Bootstrap 5, JavaScript (ES6+), Leaflet.js, Chart.js", "Responsive UI/UX, interactive map clustering, dynamic price decomposition graphs, and PWA offline capability."),
    ("Backend", "Python 3.11+, Django 5.1, Django REST Framework (DRF)", "MVC architecture, secure authentication, transactional order pipelines, RESTful APIs, and business domain logic."),
    ("Database", "SQLite (Evaluation) / PostgreSQL (Production)", "Relational schema, ACID transaction guarantees, index-optimized spatial and catalog queries."),
    ("Key Integrations", "data.gov.in Agmarknet API, OpenStreetMap, Scikit-learn, Web Audio IVR", "Official APMC mandi feeds, geospatial routing, 7-day demand forecasting regressor, and browser telephony simulation.")
]

for r_idx, (layer, stack, role) in enumerate(tech_rows, start=1):
    for c_idx, val in enumerate([layer, stack, role]):
        c = t_tech.cell(r_idx, c_idx)
        c.width = t_widths[c_idx]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        r = p.add_run(val)
        if c_idx == 0:
            r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(8)

for row in t_tech.rows:
    for c in row.cells:
        set_cell_border(c, top=dict(sz=4, val='single', color='555555'),
                           bottom=dict(sz=4, val='single', color='555555'),
                           left=dict(sz=4, val='single', color='555555'),
                           right=dict(sz=4, val='single', color='555555'))

# --- PAGE 3 ---
doc.add_page_break()

table_box3 = doc.add_table(rows=1, cols=1)
table_box3.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_box3 = table_box3.cell(0, 0)
cell_box3.width = Inches(7.0)
set_cell_border(cell_box3, top=dict(sz=8, val='single', color='222222'),
                           bottom=dict(sz=8, val='single', color='222222'),
                           left=dict(sz=8, val='single', color='222222'),
                           right=dict(sz=8, val='single', color='222222'))

p3_title = cell_box3.paragraphs[0]
p3_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
p3_title.paragraph_format.space_after = Pt(3)
r3 = p3_title.add_run("Narrative Summary of Project Status:")
r3.bold = True
r3.font.name = 'Times New Roman'
r3.font.size = Pt(10.5)

add_box_p(cell_box3, " The AgriConnect platform has successfully transitioned from conceptual system blueprint to a fully implemented, production-tested software system. All primary architectural layers—encompassing the 5-role authentication engine, assisted-digital VDF kiosk, DTMF/IVR telephony simulator, direct consumer marketplace with dynamic margin decomposition, order fulfillment dispatch hub, Leaflet GPS live route tracker, Agmarknet mandi market intelligence integration, and Scikit-learn demand forecasting engine—are completely developed and operational.", font_size=8.8)
add_box_p(cell_box3, " Comprehensive automated verification has been executed via test_full_suite.py across all 8 architectural modules, yielding 48 passed checks out of 48 total assertions (100% pass rate, 0 defects). The system is currently at the completed pre-deployment and evaluation stage, ready for real-world village cluster pilot trials.", font_size=8.8)

add_box_p(cell_box3, "", bold_prefix="Progress Summary:", font_size=10, space_after=2)
add_box_p(cell_box3, "", bold_prefix="1. Completed tasks / Activities", font_size=9.2, space_after=1)
add_box_p(cell_box3, " Developed custom user authentication and secure profile schemas catering to Farmers, Facilitators (VDFs), Consumers, Transporters, and Administrators with isolated dashboards and permission guards.", bold_prefix="a. 5-Role Role-Based Access Control (RBAC) System:", font_size=8.6, space_after=1.5)
add_box_p(cell_box3, " Engineered the Village Digital Facilitator portal alongside a browser-based dual-tone multi-frequency (DTMF) IVR and SMS notification engine for feature-phone and low-literacy farmers.", bold_prefix="b. Assisted-Digital Kiosk & IVR Telephony Simulator:", font_size=8.6, space_after=1.5)
add_box_p(cell_box3, " Implemented the consumer-facing fresh produce catalog featuring interactive 'Where Does Your Rupee Go?' financial breakdown cards, community Group Buying pools, and Leaflet.js farm cluster maps.", bold_prefix="c. Direct Marketplace & Dynamic Price Decomposition:", font_size=8.6, space_after=1.5)
add_box_p(cell_box3, " Built multi-item session cart, checkout, dispatch manifests, Leaflet GPS live tracking, two-party OTP delivery verification, and photo-supported dispute arbitration desks.", bold_prefix="d. End-to-End Order & Dispatch Logistics Pipeline:", font_size=8.6, space_after=1.5)
add_box_p(cell_box3, " Integrated live data.gov.in Agmarknet mandi benchmark feeds, dynamic fair-price corridor computation, and trained Random Forest / Ridge ML models for 7-day crop demand prediction.", bold_prefix="e. Market Intelligence & Scikit-Learn ML Forecasting:", font_size=8.6, space_after=2)

add_box_p(cell_box3, "", bold_prefix="2. Planned tasks / Activities", font_size=9.2, space_after=1)
add_box_p(cell_box3, " Execute field deployment across selected village panchayats in Jharkhand to evaluate real-world kiosk operations with active farmer and VDF cohorts.", bold_prefix="a. Live Rural Village Pilot Deployment:", font_size=8.6, space_after=1.5)
add_box_p(cell_box3, " Integrate live telecom SMS gateway (Twilio / SMSCountry) and IVR PRI cloud trunks to transition from browser simulation to live mobile network lines.", bold_prefix="b. Telecommunication Cloud Gateway Binding:", font_size=8.6, space_after=1.5)
add_box_p(cell_box3, " Transition from local SQLite evaluation environment to high-availability cloud architecture (PostgreSQL, Redis caching, Gunicorn/Nginx on AWS/GCP).", bold_prefix="c. Production Cloud Infrastructure Migration:", font_size=8.6, space_after=1.5)
add_box_p(cell_box3, " Extend automated voice listing engine to support local spoken dialects (Hindi, Santhali, Mundari, Ho) via advanced ASR speech models.", bold_prefix="d. Regional Vernacular Voice NLP Expansion:", font_size=8.6, space_after=1.5)
add_box_p(cell_box3, " Interface physical microcontroller hardware (ESP32/DHT22) for real-time transit vehicle temperature, humidity, and cold-chain logging.", bold_prefix="e. Hardware IoT Sensor Telemetry Integration:", font_size=8.6, space_after=2)

add_box_p(cell_box3, "", bold_prefix="Problems Faced (if any):", font_size=10, space_after=1.5)
add_box_p(cell_box3, " Smallholder farmers often lack smartphones and data connectivity. Resolution: Conceived and engineered the 'Phygital' Village Digital Facilitator (VDF) assisted model and dual-tone IVR phone intake.", bold_prefix="1. Rural Digital Illiteracy & Device Constraints:", font_size=8.5, space_after=1.5)
add_box_p(cell_box3, " data.gov.in Agmarknet endpoints occasionally exhibited variable response latency. Resolution: Implemented an intelligent asynchronous caching layer with verified fallback benchmark prices.", bold_prefix="2. Government Mandi API Schema Variance & Latency:", font_size=8.5, space_after=1.5)
add_box_p(cell_box3, " Breaking down prices transparently across logistics, commissions, and farmer base values required strict validation. Resolution: Engineered real-time price decomposition algorithms in the product model.", bold_prefix="3. Dynamic Price Decomposition Mathematical Balancing:", font_size=8.5, space_after=1.5)
add_box_p(cell_box3, " Transit damages risk multi-stakeholder disputes. Resolution: Developed two-party OTP delivery confirmation and photo-documented dispute arbitration workflows.", bold_prefix="4. Perishable Cold-Chain & Quality Verification:", font_size=8.5, space_after=2)

add_box_p(cell_box3, "", bold_prefix="Project Completion Percentage:", font_size=10, space_after=1.5)
add_box_p(cell_box3, " 92% COMPLETED. Evaluation Rationale: 100% of all software requirements, database models, business logic controllers, interactive UI views, and the 48-unit automated test suite have been built and rigorously verified. The remaining 8% is designated for real-world telecom PRI line binding, physical IoT sensor hookup, and rural field pilot trials.", bold_prefix="Assessment & Justification:", font_size=8.8)

# --- PAGE 4 ---
doc.add_page_break()

for i in range(1, 6):
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_after = Pt(10)
    r_sig = p_sig.add_run(f"Signature of Student {i}- ")
    r_sig.font.name = 'Times New Roman'
    r_sig.font.size = Pt(10)
    r_line = p_sig.add_run("__________________________________________________________________")
    r_line.font.name = 'Times New Roman'
    r_line.font.size = Pt(10)
    r_line.font.color.rgb = RGBColor(120, 120, 120)

def add_remarks_section_docx(title, num_lines=6, sig_label=""):
    p_rem = doc.add_paragraph()
    p_rem.paragraph_format.space_before = Pt(10)
    p_rem.paragraph_format.space_after = Pt(3)
    r_rem = p_rem.add_run(title)
    r_rem.bold = True
    r_rem.font.name = 'Times New Roman'
    r_rem.font.size = Pt(10.5)
    
    for _ in range(num_lines):
        p_l = doc.add_paragraph()
        p_l.paragraph_format.space_after = Pt(7)
        r_l = p_l.add_run("_________________________________________________________________________________")
        r_l.font.color.rgb = RGBColor(140, 140, 140)
        r_l.font.size = Pt(9)
    
    p_f = doc.add_paragraph()
    p_f.paragraph_format.space_after = Pt(12)
    r_date = p_f.add_run("Date : __________________________")
    r_date.bold = True
    r_date.font.size = Pt(9.5)
    r_space = p_f.add_run("                                        ")
    r_sig_line = p_f.add_run(f"{sig_label}")
    r_sig_line.bold = True
    r_sig_line.font.size = Pt(9.5)

add_remarks_section_docx("Project Supervisor’s Remarks:", num_lines=7, sig_label="Project Supervisor Signature")
add_remarks_section_docx("Program Coordinator’s Remarks", num_lines=6, sig_label="Program Coordinator’s Signature")
add_remarks_section_docx("Dean Remarks", num_lines=6, sig_label="Dean’s Signature:")

doc.save(DOCX_OUTPUT)
print(f"Generated DOCX report: {DOCX_OUTPUT}")
